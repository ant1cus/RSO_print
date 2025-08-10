import traceback

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pathlib import Path



def create_account_number(incoming_data: dict, current_progress: int, now_doc: int, all_doc: int, line_doing,
                          line_progress, progress_value, event, window_check, info_value) -> dict:
    logging = incoming_data['logging']
    try:
        logging.info("Создаём файл учетных номеров")
        line_doing.emit(f"Создаём файл учетных номеров {incoming_data['number_start']}")
        wb = Workbook()  # Открываем книгу
        ws = wb.active  # Активный лист
        i, j = 0, 1
        ws.column_dimensions[get_column_letter(j)].width = 12  # Ширина столбцов для отображения
        for el in range(1, 25001):  # Заполняем
            i += 1
            ws.cell(i, j).value = f"{incoming_data['number_start']}/{el}"
            if i % 100 == 0:  # Переходим на следующий столбец, чтобы немного значений в одном
                i = 0
                j += 1
                ws.column_dimensions[get_column_letter(j)].width = 12  # Ширина
        wb.save(str(Path(incoming_data['start_path'],
                         f"Файл учетных номеров № {incoming_data['number_start']}.xlsx")))
        wb.close()  # Закрываем
        line_progress.emit(f'Выполнено {int(100)} %')
        progress_value.emit(int(100))
        return {'status': 'success', 'trace': '', 'text': f"Создание файла учетных номеров успешно завершено"}
    except BaseException as exception:
        return {'status': 'error', 'trace': traceback.format_exc(),
                'text': f'Ошибка при создании файла учетных номеров - {exception}'}