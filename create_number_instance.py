import os
import shutil
import traceback

import docx
import re

from pathlib import Path

from docx.shared import Pt


def create_number_instance(incoming_data: dict, current_progress: int, now_doc: int, all_doc: int, line_doing,
                           line_progress, progress_value, event, window_check, info_value) -> dict:
    try:
        logging = incoming_data['logging']
        percent = 100/incoming_data['all_doc']
        for number_folder in incoming_data['number_instance']:
            finish_path = Path(incoming_data['finish_path'], str(number_folder) + ' экземпляр')
            if not Path(finish_path).exists():
                os.mkdir(finish_path)
            for doc in Path(incoming_data['start_path']).rglob('*.*'):
                event.wait()
                if window_check.stop_threading:
                    return {'status': 'cancel', 'trace': '', 'text': ''}
                if doc.suffix != '.docx':
                    continue
                not_need_doc = True
                for f in ['заключение', 'протокол', 'предписание']:
                    if f in doc.name.lower():
                        not_need_doc = False
                        break
                if not_need_doc:
                    continue
                if Path(finish_path, doc.name).exists():
                    logging.warning(f'В папке {finish_path} уже был документ {doc.name}')
                    continue
                line_doing.emit(f'Создаём {number_folder} экземпляр для {doc.name} ({now_doc} из {all_doc})')
                logging.info(f'Создаём {number_folder} экземпляр для {doc.name}')
                shutil.copy2(str(doc), str(Path(finish_path, doc.name)))
                doc_2 = docx.Document(str(Path(finish_path, doc.name)))
                for p_2 in doc_2.sections[0].first_page_header.paragraphs:
                    if re.findall(r'№1', p_2.text):
                        text = re.sub(r'№1', '№' + str(number_folder), p_2.text)
                        p_2.text = text
                        for run in p_2.runs:
                            run.font.size = Pt(11)
                            run.font.name = 'Times New Roman'
                        break
                for p_2 in doc_2.sections[len(doc_2.sections) - 1].first_page_footer.paragraphs:
                    if re.findall(r'Отп. 1 экз. в адрес', p_2.text):
                        text = re.sub(r'Отп. 1 экз. в адрес', 'Отп. ' + str(number_folder) + ' экз. в адрес',
                                      p_2.text)
                        p_2.text = text
                        for run in p_2.runs:
                            run.font.size = Pt(11)
                            run.font.name = 'Times New Roman'
                        break
                doc_2.save(str(Path(finish_path, doc.name)))  # Сохраняем
                current_progress += percent
                line_progress.emit(f'Выполнено {int(current_progress)} %')
                progress_value.emit(int(current_progress))
                now_doc += 1
        return {'status': 'success', 'trace': '', 'text': f"Создание экземпляров документов для "
                                                          f"{Path(incoming_data['start_path']).name} успешно завершено"}
    except BaseException as exception:
        return {'status': 'error', 'trace': traceback.format_exc(),
                'text': f'Ошибка при создании экземпляра документов - {exception}'}