import logging
import re
import shutil
import time
import traceback
import itertools
import zipfile

import docx
import pythoncom
import fitz
import os
import pandas as pd
import numpy as np
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

from word2pdf import word2pdf
from zipfile import ZipFile
from pathlib import Path
from natsort import natsorted
from typing import Any


def sorting_files(path: Path, path_new: Path, fso: bool, inventory: list) -> dict:
    """Сортирует файлы для подсчета значения прогресса и последующей вставки текстовки и номеров"""
    try:
        application_dict = {}  # Для подсчёта листов - новое.
        error_text = []
        data = {}
        if fso:
            docs = {}
            for folder in os.listdir(path):
                if 'проверке' in folder.lower():
                    if folder != 'Материалы по специальной проверке технических средств':
                        error_text.append('Название папки «Материалы по специальной проверке'
                                          ' технических средств» написано с ошибками')
                    else:
                        file = os.listdir(Path(path, folder))
                        file = natsorted(file, key=lambda y: y.rpartition(' ')[2][:-5])
                        for name_element in ['акт', 'заключение']:
                            for element in file:
                                if name_element in element.lower():
                                    docs[element] = Path(path, folder)
                elif 'исследованиям' in folder.lower():
                    if folder != 'Материалы по специальным исследованиям технических средств':
                        error_text.append('Название папки «Материалы по специальным исследованиям'
                                          ' технических средств» написано с ошибками')
                    else:
                        file = os.listdir(Path(path, folder))
                        file = natsorted(file, key=lambda y: y.rpartition(' ')[2][:-5])
                        for name_element in ['протокол', 'предписание']:
                            for element in file:
                                if name_element in element.lower():
                                    docs[element] = Path(path, folder)
                elif 'дополнительные' in folder.lower() and os.path.isdir(folder):
                    if folder != 'Дополнительные материалы':
                        error_text.append('Название папки «Дополнительные материалы» написано с ошибками')
                    else:
                        shutil.copytree(Path(path, folder), Path(path_new, folder))
            if error_text:
                return {'error': True, 'text': error_text, 'data': {}}
            for element in os.listdir():
                if 'сопроводит' in element.lower():
                    docs[element] = path
            data['docx_for_progress'] = len(docs)
        else:
            docs = [file for file in os.listdir(path) if file.endswith('.docx')]  # Список документов
            inventory_list = [file.name for file in inventory]
            docs = [*docs, *inventory_list]
            docs = natsorted(docs, key=lambda y: y.rpartition(' ')[2][:-5])
            docs_ = [j_ for i_ in
                     ['^Акт', 'Приложение \d? к акту', '^Заключение', 'Приложение \d? к заключению', 'Протокол',
                      'Приложение А', 'Предписание', 'Форма 3', 'Опись', 'Сопроводит'] for j_ in docs if
                     re.findall(i_, j_, re.I)]
            docs_not = [i_ for i_ in docs if i_ not in docs_ and '~' not in i_]
            docs = docs_not + docs_
            # Процент для прогресса
            docx_for_progress = 0
            for name_file in os.listdir(path):
                if re.findall(r'приложение а', name_file.lower()):
                    # Удалить это отсюда, не нужно
                    with ZipFile(Path(path, name_file)) as my_doc:
                        xml_content = my_doc.read('docProps/app.xml')  # Общие свойства
                        pages = int(re.findall(r'<Pages>(\w*)</Pages>', xml_content.decode())[0])
                    if pages == 1:
                        pythoncom.CoInitializeEx(0)
                        # self.logging.info(f"Считаем кол-во листов в приложении {name_file}")
                        # status.emit(f"Считаем кол-во листов в приложении {name_file}")
                        word2pdf(str(Path(path, name_file)), str(Path(path, name_file + '.pdf')))
                        input_file = fitz.open(str(Path(path, name_file + '.pdf')))  # Открываем
                        pages = input_file.page_count  # Получаем кол-во страниц
                        input_file.close()  # Закрываем
                        os.remove(str(Path(path, name_file + '.pdf')))  # Удаляем pdf документ
                    application_dict[name_file] = pages
                else:
                    docx_for_progress += 1
            data['docs_for_progress'] = docx_for_progress
        data['docs'] = [Path(path, doc) for doc in docs]
        return {'error': False, 'text': '', 'data': data}
    except BaseException as exception:
        return {'error': True, 'text': str(exception), 'trace': traceback.format_exc()}


def report_rso(path: Path):
    """Функция для генерации отчёта МВД"""
    file_mo = [mo for mo in os.listdir(path) if mo[-3:] == 'txt' and 'F19' in mo][0]
    df_report_rso = pd.read_csv(Path(path, file_mo), delimiter='|', encoding='ANSI', names=[
        'Порядковый номер лицензиата',
        'Серийный номер комплекта',
        'Серийный номер системного блока', 'удалить'])
    df_report_rso['Порядковый номер лицензиата'] = df_report_rso['Порядковый номер лицензиата'].astype(str)
    df_report_rso['№'] = np.arange(1, 1 + len(df_report_rso))
    df_report_rso = df_report_rso.reindex(columns=['№',
                                                   'Порядковый номер лицензиата',
                                                   'Серийный номер комплекта',
                                                   'Серийный номер системного блока',
                                                   'Заключение',
                                                   'Кол-во листов закл.',
                                                   'Протокол',
                                                   'Кол-во листов прот.',
                                                   'Предписание',
                                                   'Кол-во листов пред.',
                                                   'Сумма листов на комплект'])


def sp_sorting(name_gk: str, finish_path: Path, sp_path_dir: Path, sp_path_file: Path, check_sp: list,
               progress_value, line_progress) -> dict:
    errors = []
    try:
        if name_gk:
            logging.info("Создаём папку для СП, если её нет")
            path_dir_sp = Path(sp_path_dir, name_gk)
            os.makedirs(path_dir_sp, exist_ok=True)
        else:
            path_dir_sp = Path(sp_path_dir)
        df_number_sp = pd.read_excel(str(Path(sp_path_file)), sheet_name=0, header=None)
        df_number_sp.fillna(False, inplace=True)
        df_number_sp.drop(0, inplace=True)
        for name1, name2 in itertools.zip_longest(df_number_sp[0], df_number_sp[1]):
            if name1:
                os.makedirs(str(Path(path_dir_sp, str(name1) + ' В')), exist_ok=True)
            if name2:
                os.makedirs(str(Path(path_dir_sp, str(name2))), exist_ok=True)
        files = [j_ for i_ in ['акт', 'заключение', 'протокол', 'предписание', 'инфокарта', 'result']
                 for j_ in os.listdir(finish_path) if re.findall(i_.lower(), j_.lower())]
        percent = 5 / len(files)
        current_progress = 90
        for file in files:
            if 'акт' in file.lower() or 'result' in file.lower():
                shutil.copy(str(Path(finish_path, file)), str(Path(path_dir_sp)))
            else:
                no_sn_in_sp = True
                sn_number = file.rpartition(' ')[0].rpartition(' ')[2]
                for folder_sp in os.listdir(str(path_dir_sp)):
                    if re.findall(sn_number, folder_sp):
                        no_sn_in_sp = False
                        shutil.copy(str(Path(finish_path, file)), str(Path(path_dir_sp, folder_sp)))
                if no_sn_in_sp:
                    errors.append('Документ с с.н. ' + sn_number + ' (' + file + ') не найден в материалах СП')
            current_progress += percent
            line_progress.emit(f'Выполнено {int(current_progress)} %')
            progress_value.emit(int(current_progress))
        for folder_sp in os.listdir(str(Path(path_dir_sp))):
            if os.path.isdir(str(Path(path_dir_sp, folder_sp))):
                file_sp = [file.partition(' ')[0].lower() for file in
                           os.listdir(str(Path(path_dir_sp, folder_sp)))]
                for ind, check_file in enumerate(['заключение', 'протокол', 'предписание', 'инфокарта']):
                    if check_sp[ind] and (check_file in file_sp) is False:
                        errors.append('В папке ' + str(folder_sp) + ' отсутствует ' + check_file)
                current_progress += percent
                line_progress.emit(f'Выполнено {int(current_progress)} %')
                progress_value.emit(int(current_progress))
        if errors:
            return {'status': 'warning', 'trace': '', 'text': errors}
        return {'status': 'success', 'trace': '', 'text': ''}
    except BaseException as es:
        return {'status': 'error', 'trace': traceback.format_exc(), 'text': f'Ошибка при сортировки СП - {es}',
                'data': ''}


def pages_count(file: Path, minus: bool = False) -> dict:
    """Функция для подсчёта количества страниц в документе. Принимает путь к файлу. Переменная minus служить для
    вычитания одной страницы (фонаря), если он уже присутствует в документе. Такое бывает, когда листы считаются в
    динамически заполняемом документе (сопровод или опись, например)"""
    try:
        name = file.name
        parent_path = file.parent
        pythoncom.CoInitializeEx(0)
        name_pdf = name + '.pdf'
        word2pdf(str(Path(parent_path, name)), str(Path(parent_path, name_pdf)))
        input_file_pdf = fitz.open(str(Path(parent_path, name_pdf)))  # Открываем пдф
        count_page = input_file_pdf.page_count  # Получаем кол-во страниц
        input_file_pdf.close()  # Закрываем
        count_page = count_page - 1 if minus else count_page
        os.remove(str(Path(parent_path, name_pdf)))  # Удаляем пдф документ
        temp_docx = os.path.join(parent_path, name)
        temp_zip = os.path.join(parent_path, name + ".zip")
        temp_folder = os.path.join(parent_path, "template")

        if os.path.exists(temp_zip):
            rm(temp_zip)
        if os.path.exists(temp_folder):
            rm(temp_folder)
        if os.path.exists(Path(parent_path, 'zip')):
            rm(Path(parent_path, 'zip'))
        os.rename(temp_docx, temp_zip)
        os.mkdir(Path(parent_path, 'zip'))
        with ZipFile(temp_zip) as my_document:
            my_document.extractall(temp_folder)
        pages_xml = os.path.join(temp_folder, "docProps", "app.xml")
        string = open(pages_xml, 'r', encoding='utf-8').read()
        string = re.sub(r"<Pages>(\w*)</Pages>", "<Pages>" + str(count_page) + "</Pages>", string)
        with open(pages_xml, "wb") as file_wb:
            file_wb.write(string.encode("UTF-8"))
        try_number = 0
        while True:
            try:
                os.remove(temp_zip)
                break
            except PermissionError:
                if try_number == 4:
                    break
                time.sleep(3)
                try_number += 1
        if try_number == 4:
            return {'error': True, 'text': 'Не удалось удалить файл'}
        shutil.make_archive(temp_zip.replace(".zip", ""), 'zip', temp_folder)
        os.rename(temp_zip, temp_docx)  # rename zip file to docx
        rm(temp_folder)
        rm(Path(parent_path, 'zip'))
        return {'error': False, 'text': '', 'pages': count_page, 'trace': ''}
    except BaseException as exception:
        return {'error': True, 'text': exception, 'pages': 0, 'trace': traceback.format_exc()}


def rm(folder_path):
    try:
        while len(os.listdir(folder_path)) != 0:
            time.sleep(0.5)
            for file_object in os.listdir(folder_path):
                flag_ = True
                while flag_:
                    try:
                        file_object_path = os.path.join(folder_path, file_object)
                        try:
                            if os.path.isfile(file_object_path) or os.path.islink(file_object_path):
                                os.remove(file_object_path)
                            else:
                                try:
                                    shutil.rmtree(file_object_path)
                                except FileNotFoundError:
                                    pass
                        except OSError:
                            os.remove(folder_path)
                        flag_ = False
                    except BaseException:
                        pass
    except NotADirectoryError:
        os.remove(folder_path)
    time.sleep(0.05)
    shutil.rmtree(folder_path)


def inventory_insert(documents: pd, incoming_data: dict, incoming_path: Path) -> dict:
    """Добавление номеров и текстовок для описей"""

    def df_add_inventory(number, finish_path, docs, ind):
        docs.loc[ind, 'name'] = f'Опись №{number}.docx'
        docs.loc[ind, 'number'] = number
        docs.loc[ind, 'finish_path'] = Path(finish_path, f'Опись №{number}.docx')
        docs.loc[ind, 'parent_path'] = finish_path
        docs.loc[ind, 'change_date'] = True
        docs.loc[ind, 'classified'] = incoming_data['classified']
        docs.loc[ind, 'executor'] = incoming_data['inventory_executor']
        docs.loc[ind, 'first_header_text'] = f'{incoming_data["classified"]}\n(без приложения не секретно)\nЭкз.№ 1'
        docs.loc[ind, 'text'] = '\n\n' + incoming_data["account_position"] + '\t\t\t\t\t\t\t\t' + incoming_data["account_executor"]
        docs.loc[ind, 'date'] = incoming_data['date']
        return docs

    try:
        number_inventory = 1
        if incoming_data['flag_inventory'] == 1:
            index = documents.loc[documents['name'].str.contains(f'Опись №{number_inventory}.docx', case=False)].index[0]
            documents = df_add_inventory(number_inventory, incoming_path, documents, index)
            documents.loc[documents['account_list'] == 1, 'account_list'] = f'Опись №{number_inventory}.docx'
        else:
            conclusion_num = len(documents[(documents['name'].str.contains(r'заключение', case=False))])
            conclusions = documents[(documents['name'].str.contains(r'заключение', case=False))].reset_index()
            prescriptions = documents[(documents['name'].str.contains(r'предписание', case=False))].reset_index()
            # conclusion_num = len(documents[(documents['name'].str.contains(r'заключение', case=False)
            # 								& (documents['parent_path'] == incoming_path))])
            while True:
                index = documents.loc[documents['name'].str.contains(f'Опись №{number_inventory}.docx', case=False)].index[0]
                documents = df_add_inventory(number_inventory, incoming_path, documents, index)
                docs_40 = []
                for doc in [conclusions, prescriptions]:
                    slice_docs = doc.iloc[1:41, 'index'].tolist()
                    documents.loc[slice_docs, 'account_list'] = f'Опись №{number_inventory}.docx'
                    docs_40.append(doc.iloc[40:])
                conclusions, prescriptions = docs_40[0], docs_40[1]
                if conclusion_num - 40 <= 0:
                    break
                else:
                    conclusion_num = conclusion_num - 40
                    number_inventory += 1
        return {'status': 'success', 'trace': '', 'text': '', 'data': documents}
    except BaseException as es:
        return {'status': 'error', 'trace': traceback.format_exc(), 'text': f'Ошибка при добавлении описи(ей) - {es}',
                'data': ''}


def sort_order(name: str) -> int:
    if 'акт' in name.lower():
        return 1
    elif 'заключение' in name.lower():
        return 2
    elif 'протокол' in name.lower():
        return 3
    elif 'приложение а' in name.lower():
        return 4
    elif 'предписание' in name.lower():
        return 5
    else:
        return 99


def list_doc(dp):
    # Открываем word документ как зип архив для доступа к xml свойствам
    with zipfile.ZipFile(dp) as my_doc:
        xml_content = my_doc.read('docProps/app.xml')  # Общие свойства
        pages_ = re.findall(r'<Pages>(\w*)</Pages>', xml_content.decode())  # Ищем кол-во страниц
        ns = int(pages_[0])
    return ns


def delete_header_footer_second_acc(path: Path, text_first_header: str, secret_num: str, text_for_foot: str,
                                    para: bool) -> dict:
    try:
        doc = docx.Document(str(path))  # Открываем

        def paragraph_del(par):
            par.text = None
            paragraph_ = par._element
            paragraph_.getparent().remove(paragraph_)
            paragraph_._p = paragraph_._element = None

        list_paragraph = []
        for enum, paragraph in enumerate(doc.sections[0].first_page_header.paragraphs):
            if 'экз.' in paragraph.text.lower():
                list_paragraph = [i for i in range(enum + 1)]
                break
        for paragraph in list_paragraph:
            doc.sections[0].first_page_header.paragraphs[paragraph].text = None
        for paragraph in range(len(list_paragraph) - 1):
            p = doc.sections[0].first_page_header.paragraphs[paragraph]._element
            p.getparent().remove(p)
            p._p = p._element = None
        p = doc.sections[0].first_page_header.paragraphs[0]._element
        p.getparent().remove(p)
        p._p = p._element = None
        doc.sections[0].first_page_footer.paragraphs[0].text = None
        if len(doc.sections) == 1:
            for paragraph in doc.sections[len(doc.sections) - 1].footer.paragraphs:
                if 'Б/ч' in paragraph.text:
                    paragraph_del(paragraph)
                    break
                paragraph_del(paragraph)
        else:
            doc.sections[0].footer.paragraphs[0].text = None
            if doc.sections[len(doc.sections) - 1].different_first_page_header_footer:
                for paragraph in doc.sections[len(doc.sections) - 1].first_page_footer.paragraphs:
                    if 'Б/ч' in paragraph.text:
                        paragraph_del(paragraph)
                        break
                    paragraph_del(paragraph)
                doc.sections[
                    len(doc.sections) - 1].first_page_header.is_linked_to_previous = False  # header
                doc.sections[len(doc.sections) - 1].first_page_footer.is_linked_to_previous = False  # Футер
            else:
                for paragraph in doc.sections[len(doc.sections) - 1].footer.paragraphs:
                    if 'Б/ч' in paragraph.text:
                        paragraph_del(paragraph)
                        break
                    paragraph_del(paragraph)
        while True:
            flag_for_exit = 0
            if flag_for_exit == 3:
                break
            try:
                doc.save(str(path))
                break
            except PermissionError:
                flag_for_exit += 1
                time.sleep(3)
        doc = docx.Document(str(path))
        sectPrs = doc._element.xpath(".//w:pPr/w:sectPr")
        for sectPr in sectPrs:
            sectPr.getparent().remove(sectPr)
        doc.add_section()  # Добавляем последнюю страницу
        if para:
            last = doc.sections[
                len(doc.sections) - 1].first_page_header  # Колонтитул для последней страницы
            last.is_linked_to_previous = False  # Отвязываем от предыдущей секции чтобы не повторялись
            foot = doc.sections[len(doc.sections) - 1].first_page_footer  # Нижний колонтитул
            foot.is_linked_to_previous = False  # Отвязываем
        else:
            last = doc.sections[len(doc.sections) - 1].header  # Колонтитул для последней страницы
            last.is_linked_to_previous = False  # Отвязываем от предыдущей секции чтобы не повторялись
            foot = doc.sections[len(doc.sections) - 1].footer  # Нижний колонтитул
            foot.is_linked_to_previous = False  # Отвязываем
        if doc.sections[0].different_first_page_header_footer:
            header = doc.sections[0].first_page_header  # Верхний колонтитул первой страницы
            footer = doc.sections[0].first_page_footer
        else:
            header = doc.sections[0].header
            footer = doc.sections[0].footer
        if len(footer.paragraphs) == 0:
            footer.add_paragraph()
        footer.paragraphs[0].text = secret_num
        head = header.paragraphs[0]  # Параграф
        head.insert_paragraph_before(text_first_header)  # Вставляем перед колонтитулом
        head = header.paragraphs[0]  # Выбираем новый первый параграф
        head_format = head.paragraph_format  # Настройки параграфа
        head_format.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT  # Выравниваем по правому краю
        # Текст для фонарика
        foot.paragraphs[0].text = text_for_foot
        doc.save(str(path))
        return {'status': 'success', 'trace': '', 'text': f"Документ {path.name} успешно сохранён"}
    except BaseException as es:
        return {'status': 'error', 'trace': traceback.format_exc(),
                'text': f'Ошибка при удалении шапки в сопроводе - {es}'}


def return_error(log: logging, warning: str, status, status_text: str, default_path: Path, status_finish: list,
                 window, event: Any = False, error: str = '', info_value=None) -> None:
    """Функция для возврата ошибок по всей программе и отмены оперции пользователем. """
    if info_value is None:
        info_value = []
    if event:
        log.error(error)
        if isinstance(info_value[2], list):
            error_text = '\n'.join(info_value[2])
        else:
            error_text = 'Работа программы завершена из-за непредвиденной ошибки, обратитесь к разработчику'
        info_value[0].emit(info_value[1], error_text)
        event.clear()
        event.wait()
        log.error(warning)
        status.emit(status_text)
        os.chdir(default_path)
        status_finish[0].emit(status_finish[1], status_finish[2])
        time.sleep(1)  # Не удалять, не успевает отработать emit status_finish. Может потом
        window.close()
        return
