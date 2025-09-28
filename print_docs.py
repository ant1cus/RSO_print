import datetime
import os
import re
import time
import traceback
from pathlib import Path
import docx
import fitz
import numpy as np
import pythoncom
import win32com
import win32print
import getpass
import socket

from PyQt5 import QtPrintSupport
from natsort import natsorted
from openpyxl import load_workbook
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement, ns
from small_functions import print_doc

import pandas as pd

from word2pdf import word2pdf


def del_pdf_files(logging, del_path):
    logging.info("Удаляем все созданные pdf из папки")
    pdf_files = [i for i in os.listdir(del_path) if i.endswith('.pdf')]
    if pdf_files:
        logging.info(f"Удаляем пдф после ошибки")
        for pdf_file in pdf_files:
            try:
                os.remove(str(Path(del_path, pdf_file)))
            except Exception:
                time.sleep(3)
                input_file = fitz.open(str(Path(del_path, pdf_file)))
                input_file.close()
                try:
                    os.remove(str(Path(del_path, pdf_file)))
                except Exception as e:
                    logging.error(f"Не удалось удалить файл {Path(del_path, pdf_file)}")
                    logging.error(f"Ошибка:\n {e}'\n'{traceback.format_exc()}")


def create_element(attrib_name):
    return OxmlElement(attrib_name)


def create_attribute(attrib, attrib_name, attrib_value):
    attrib.set(ns.qn(attrib_name), attrib_value)


def add_page_number(paragraph, value_num, number_page=''):
    page_run = paragraph.add_run()
    t1 = create_element('w:t')
    create_attribute(t1, 'xml:space', 'preserve')
    t1.text = '\t\t' + value_num
    page_run._r.append(t1)

    page_num_run = paragraph.add_run()

    fld_char1 = create_element('w:fldChar')
    create_attribute(fld_char1, 'w:fldCharType', 'begin')

    instr_text_or1 = create_element('w:instrText')
    create_attribute(instr_text_or1, 'xml:space', 'preserve')
    instr_text_or1.text = "="

    fld_char2 = create_element('w:fldChar')
    create_attribute(fld_char2, 'w:fldCharType', 'begin')

    instrText = create_element('w:instrText')
    create_attribute(instrText, 'xml:space', 'preserve')
    instrText.text = "PAGE"

    fld_char3 = create_element('w:fldChar')
    create_attribute(fld_char3, 'w:fldCharType', 'end')

    instr_text_or2 = create_element('w:instrText')
    create_attribute(instr_text_or2, 'xml:space', 'preserve')
    instr_text_or2.text = " - 2 +" + number_page

    fld_char4 = create_element('w:fldChar')
    create_attribute(fld_char4, 'w:fldCharType', 'end')

    page_num_run._r.append(fld_char1)
    page_num_run._r.append(instr_text_or1)
    page_num_run._r.append(fld_char2)
    page_num_run._r.append(instrText)
    page_num_run._r.append(fld_char3)
    page_num_run._r.append(instr_text_or2)
    page_num_run._r.append(fld_char4)


def folder_print(incoming_data: dict, start_path: Path, line_doing, line_progress, progress_value, event, window_check,
                 info_value) -> dict:
    logging = incoming_data['logging']
    current_progress = 0
    try:
        line_doing.emit(f'Готовим печать документов в «{start_path.name}»')
        errors = []
        columns_name = ['name', 'start_path', 'parent_path', 'print_order', 'print', 'print_num', 'pages', 'pdf_path']
        documents = pd.DataFrame(columns=columns_name)
        # Проверка на количество листов и учетных номеров
        docs = [i for i in os.listdir(start_path) if i[-4:] == 'docx' and '~' not in i]  # Список файлов
        logging.info('Сортируем')
        docs = natsorted(docs, key=lambda y: y.rpartition(' ')[2][:-5])
        for index, file in enumerate(Path(start_path).rglob('*.*')):
            documents = pd.concat([documents, pd.DataFrame({'name': [file.name], 'start_path': [file],
                                                            'parent_path': [file.parent], 'print_order': [index],
                                                            'print': [False], 'print_num': [''],
                                                            'pdf_path': [Path(file.parent, file.stem + '.pdf')]
                                                            })], ignore_index=True)
        if incoming_data['print_order']:
            logging.info('Есть порядок печати')
            quantity_docs = {'Заключение': 0, 'Протокол': 0, 'Приложение А': 0, 'Предписание': 0}
            docs_name = {}
            for element in docs:
                for doc_name in ['Заключение', 'Протокол', 'Приложение А', 'Предписание']:
                    if re.findall(doc_name.lower(), element.lower()):
                        quantity_docs[doc_name] += 1
                        doc_number = element.rpartition('.')[0].rpartition(' ')[2]
                        if doc_number in docs_name:
                            docs_name[doc_number].append(element)
                        else:
                            docs_name[doc_number] = [element]
            docs_ = []
            for element in docs_name:
                for doc_name in ['Заключение', 'Протокол', 'Приложение А', 'Предписание']:
                    for i in docs_name[element]:
                        if re.findall(doc_name.lower(), i.lower()):
                            docs_.append(i)
                            break
            docs_sec = [j for i in ['Форма 3', 'Опись',
                                    'Сопровод'] for j in docs if re.findall(i.lower(), j.lower())]
            docs_ = docs_ + docs_sec
        else:
            logging.info('Нет порядка печати')
            docs_ = [j for i in ['Заключение', 'Протокол', 'Приложение А', 'Предписание', 'Форма 3', 'Опись',
                                 'Сопровод'] for j in docs if re.findall(i.lower(), j.lower())]
        docs_not = [i for i in docs if i not in docs_]
        docs = docs_not + docs_
        # Берем все подряд номера, потом вернём и удалим
        print_numbers_df = pd.read_excel(incoming_data['path_account_num'], header=None)
        print_numbers_df.fillna(False, inplace=True)
        print_nums = [print_numbers_df[col].to_numpy().tolist() for col in print_numbers_df.columns]
        print_nums = [x for y in print_nums for x in y if x is not False]
        print_num = 0
        one_time_load = True
        try:
            percent = 20 / len(docs)
        except ZeroDivisionError:
            return {'status': 'warning', 'text': 'Деление на 0, ни одного документа для печати', 'trace': ''}
        now_doc = 1
        all_doc = len(docs)
        for index, file in enumerate(docs):
            event.wait()
            if window_check.stop_threading:
                del_pdf_files(logging, start_path)
                return {'status': 'cancel', 'trace': '', 'text': ''}
            pythoncom.CoInitializeEx(0)
            index_doc = documents.loc[documents['name'] == file].index[0]
            if re.findall(r'заключение', file.lower(), re.I) and incoming_data['conclusion'] is False:
                documents.loc[index_doc, 'print'] = False
            elif re.findall(r'протокол', file.lower(), re.I) and incoming_data['protocol'] is False:
                documents.loc[index_doc, 'print'] = False
            elif re.findall(r'предписание', file.lower(), re.I) and incoming_data['prescription'] is False:
                documents.loc[index_doc, 'print'] = False
            elif re.findall(r'приложение а', file.lower(), re.I) and incoming_data['service']:
                documents.loc[index_doc, 'print'] = False
            else:
                documents.loc[index_doc, 'print'] = True
            documents.loc[index_doc, 'print_order'] = index
            if documents.loc[index_doc, 'print'] is False:
                current_progress += percent
                line_progress.emit(f'Выполнено {int(current_progress)} %')
                progress_value.emit(int(current_progress))
                now_doc += 1
                continue
            line_doing.emit(f'Считаем листы для печати {file} ({now_doc} из {all_doc})')
            try:
                word2pdf(str(documents.loc[index_doc, 'start_path']), str(documents.loc[index_doc, 'pdf_path']))
            except BaseException:
                word = win32com.client.Dispatch("Word.Application")
                word.Quit()
            input_file = fitz.open(str(documents.loc[index_doc, 'pdf_path']))  # Открываем пдф
            documents.loc[index_doc, 'pages'] = input_file.page_count - 1 if input_file.page_count > 1 else 1
            input_file.close()
            os.remove(str(documents.loc[index_doc, 'pdf_path']))
            current_progress += percent
            line_progress.emit(f'Выполнено {int(current_progress)} %')
            progress_value.emit(int(current_progress))
            now_doc += 1
        documents = documents.sort_values('print_order')
        documents.reset_index(drop=True, inplace=True)
        for index, pages in enumerate(documents['pages'].to_numpy().tolist()):
            if np.isnan(pages):
                continue
            if len(print_nums) < pages and one_time_load:
                if incoming_data['check_box_add_account_num'] is False:
                    errors.append('Не хватает учетных номеров, загрузите дополнительный файл')
                    break
                wb = load_workbook(incoming_data['add_path_account_num'])  # Открываем книгу
                ws = wb.active  # Делаем активный лист
                for j in range(1, ws.max_column + 1):  # По столбцам
                    for i in range(1, ws.max_row + 1):  # По строкам
                        if ws.cell(i, j).value:  # Если есть значение
                            print_nums.append(ws.cell(i, j).value)
                add_print_numbers_df = pd.read_excel(incoming_data['add_path_account_num'], header=None)
                add_print_numbers_df.fillna(False, inplace=True)
                add_print_nums = [add_print_numbers_df[col].to_numpy().tolist() for col in add_print_numbers_df.columns]
                add_print_nums = [x for y in add_print_nums for x in y if x is not False]
                print_nums = [*print_nums, *add_print_nums]
                one_time_load = False
            if len(print_nums) < pages:
                errors.append('Не хватает учетных номеров, загрузите другой файл')
                break
            documents.loc[index, 'print_num'] = '|'.join(print_nums[print_num: print_num + pages])
            print_num = print_num + pages
        now_doc = 1
        all_doc = documents['print'].sum()
        if errors:
            del_pdf_files(logging, start_path)
            return {'status': 'warning', 'text': errors, 'trace': ''}
        try:
            percent = 80 / all_doc
        except ZeroDivisionError:
            del_pdf_files(logging, start_path)
            return {'status': 'warning', 'text': 'Деление на 0, ни одного документа для печати', 'trace': ''}
        win32print.SetDefaultPrinter(incoming_data['name_printer'])
        user_name = getpass.getuser()
        printer = QtPrintSupport.QPrinterInfo.defaultPrinterName()
        computer_name = socket.gethostname()
        name_printer = incoming_data['name_printer']
        del_numbers = []
        date_for_saving = datetime.date.today()
        if not Path.exists(Path(incoming_data['default_path'], 'printing_data')):
            os.makedirs(Path(incoming_data['default_path'], 'printing_data'))
        save_printing_data_file = Path(incoming_data['default_path'], 'printing_data', str(date_for_saving) + '.txt')
        if not os.path.exists(save_printing_data_file):
            with open(save_printing_data_file, 'w'):
                pass
        # после каждой успешной печати удалять номера или записать все успешные печати и потом почистить файл
        for doc in documents.itertuples():
            event.wait()
            if window_check.stop_threading:
                del_pdf_files(logging, start_path)
                return {'status': 'cancel', 'trace': '', 'text': ''}
            try:
                if doc.print is False:
                    continue
                line_doing.emit(f'Печатаем {doc.name} ({now_doc} из {all_doc})')
                printing_date = [computer_name, user_name, str(doc.start_path), str(datetime.date.today()), printer]
                print_nums = doc.print_num.split('|')
                pythoncom.CoInitializeEx(0)
                logging.info(f"Форматируем документ {doc.name}")
                if re.findall(r'приложение а', doc.name.lower()):
                    logging.info(f"Запускаем в печать {doc.name}")
                    answer = print_doc(doc.start_path, name_printer, 1, logging, del_numbers)
                    if answer['status'] != 'success':
                        errors.append(answer['text'])
                        logging.error(answer['text'])
                    if answer['status'] == 'error':
                        logging.error(answer['trace'])
                    del_numbers = answer['data']
                else:
                    num_start = print_nums[0]
                    num_second_page = '' if len(print_nums) == 1 else print_nums[1]
                    num_stop = print_nums[-1]
                    logging.info(f"Вставляем номера листов {doc.name}")
                    form_27_data = {
                        'check_form_27': incoming_data['check_box_from_27'], 'package': incoming_data['package'],
                        'start_path': start_path, 'path_form_27': incoming_data['path_form_27'],
                        'doc_name': doc.name.lower(), 'num_start': num_start, 'num_stop': num_stop
                    }
                    word_doc = docx.Document(doc.start_path)  # Открываем
                    if word_doc.sections[0].different_first_page_header_footer:
                        footer_1 = word_doc.sections[0].first_page_footer  # Нижний колонтитул первой страницы
                    else:
                        footer_1 = word_doc.sections[0].footer  # Нижний колонтитул первой страницы
                    # footer_1 = word_doc.sections[0].first_page_footer  # Нижний колонтитул первой страницы
                    if len(footer_1.paragraphs) == 0:
                        footer_1.add_paragraph()
                    foot_1 = footer_1.paragraphs[0]  # Параграф
                    if len(footer_1.paragraphs[0].text) > 0:
                        foot_1.text = footer_1.paragraphs[0].text + '\t\t' + num_start  # Текст
                    else:
                        foot_1.text = '\t\t' + num_start  # Текст
                    foot_format = foot_1.paragraph_format  # Настройки параграфа
                    foot_format.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT  # Выравнивание по левому краю
                    if num_second_page:
                        try:
                            if len(word_doc.sections[1].footer.paragraphs) == 0:
                                word_doc.sections[1].footer.paragraphs.add_paragraph()
                            footer_2 = word_doc.sections[1].footer.paragraphs[0]  # Нижний колонтитул страницы
                        except IndexError as ex:
                            logging.warning(f"В файле {doc.name} нет второй секции!")
                            if len(word_doc.sections[0].footer.paragraphs) == 0:
                                word_doc.sections[0].footer.paragraphs.add_paragraph()
                            footer_2 = word_doc.sections[0].footer.paragraphs[0]  # Нижний колонтитул страницы
                        mask_page = num_second_page.rpartition('/')[0] + '/'
                        start_number = num_second_page.rpartition('/')[2]
                        add_page_number(footer_2, mask_page, start_number)
                    word_doc.save(doc.start_path)  # Сохраняем
                    try:
                        word2pdf(str(doc.start_path), str(doc.pdf_path))
                    except BaseException:
                        word = win32com.client.Dispatch("Word.Application")
                        word.Quit()
                    doc_old = docx.Document(doc.start_path)  # Открываем
                    last = doc_old.sections[len(doc_old.sections) - 1].first_page_footer  # Колонтитул
                    number = last.paragraphs[0].text.partition('\n')[0].rpartition(' ')[2]
                    form_27_data['number'] = number
                    if incoming_data['one_side']:  # Если печать односторонняя - печатаем
                        logging.info(f"Печатаем документ {doc.name}")
                        answer = print_doc(doc.pdf_path, name_printer, 1, logging,
                                           del_numbers, print_nums, form_27_data)
                        if answer['status'] != 'success':
                            errors.append(answer['text'])
                            logging.error(answer['text'])
                        if answer['status'] == 'error':
                            logging.error(answer['trace'])
                        del_numbers = answer['data']
                    elif incoming_data['last_duplex']:  # Если последняя страница дуплекс
                        # Дефолтный принтер
                        logging.info(f"Преобразуем документ {doc.name}")
                        input_file = fitz.open(doc.pdf_path)  # Открываем пдф
                        pages = input_file.page_count  # Получаем кол-во страниц
                        if pages == 2:
                            answer = print_doc(doc.pdf_path, name_printer, 2, logging,
                                               del_numbers, print_nums, form_27_data)
                            if answer['status'] != 'success':
                                errors.append(answer['text'])
                                logging.error(answer['text'])
                            if answer['status'] == 'error':
                                logging.error(answer['trace'])
                            del_numbers = answer['data']
                        else:
                            logging.info(f"Преобразуем документ {doc.pdf_path}")
                            output_1_side = Path(doc.pdf_path.parent, '1_' + doc.pdf_path.name)
                            output_2_side = Path(doc.pdf_path.parent, '2_' + doc.pdf_path.name)
                            # Страницы для односторонней печати
                            selected_page = [page for page in range(0, pages - 2)]
                            input_file.select(selected_page)  # Выбираем страницы
                            input_file.save(output_1_side)  # Сохраняем файл
                            # Печатаем
                            answer = print_doc(output_1_side, name_printer, 1, logging, del_numbers)
                            if answer['status'] != 'success':
                                errors.append(answer['text'])
                                logging.error(answer['text'])
                            if answer['status'] == 'error':
                                logging.error(answer['trace'])
                            del_numbers = answer['data']
                            os.remove(output_1_side)
                            input_file = fitz.open(doc.pdf_path)  # Открываем пдф
                            selected_page = [pages - 2, pages - 1]  # Страницы для двухсторонней печати
                            input_file.select(selected_page)  # Выбираем страницы
                            input_file.save(output_2_side)  # Сохраняем
                            answer = print_doc(output_2_side, name_printer, 2, logging,
                                               del_numbers, print_nums, form_27_data)
                            if answer['status'] != 'success':
                                errors.append(answer['text'])
                                logging.error(answer['text'])
                            if answer['status'] == 'error':
                                logging.error(answer['trace'])
                            del_numbers = answer['data']
                            os.remove(output_2_side)
                        input_file.close()
                    else:
                        answer = print_doc(doc.pdf_path, name_printer, 2, logging,
                                           del_numbers, print_nums, form_27_data)
                        if answer['status'] != 'success':
                            errors.append(answer['text'])
                            logging.error(answer['text'])
                        if answer['status'] == 'error':
                            logging.error(answer['trace'])
                        del_numbers = answer['data']
                    logging.info('Записываем данные с печати')
                    with open(save_printing_data_file, 'a') as f:
                        f.write(';'.join(printing_date) + '\n')
            except Exception as ex:
                logging.error("Упс, сорвалась печать файла!")
                logging.error("Ошибка:\n " + str(ex) + '\n' + traceback.format_exc())
            pdf_files = [i for i in os.listdir(doc.parent_path) if doc.name.rpartition('.')[0] in i
                         and i.endswith('.pdf')]
            if pdf_files:
                logging.info(f"Удаляем пдф после печати")
                for pdf_file in pdf_files:
                    while True:
                        permission = 0
                        try:
                            permission += 1
                            os.remove(str(Path(doc.parent_path, pdf_file)))
                            break
                        except PermissionError as per:
                            time.sleep(3)
                            logging.warning(f"Ошибка удаления файла {pdf_file} после печати - {per}")
                            if permission == 3:
                                break
            current_progress += percent
            line_progress.emit(f'Выполнено {int(current_progress)} %')
            progress_value.emit(int(current_progress))
            now_doc += 1
        logging.info(f"Удаляем напечатанные номера")
        line_doing.emit(f"Удаляем напечатанные номера")
        print_numbers_df = pd.read_excel(incoming_data['path_account_num'], header=None)
        print_numbers_df.fillna(False, inplace=True)
        print_nums = [print_numbers_df[col].to_numpy().tolist() for col in print_numbers_df.columns]
        print_nums = [x for y in print_nums for x in y if x is not False]
        final_nums = [x for x in print_nums if x not in del_numbers]
        shape_df = print_numbers_df.shape[0]
        final_nums = {enum: final_nums[i:i + shape_df] for enum, i in enumerate(range(0, len(final_nums), shape_df))}
        dict_keys = list(final_nums.keys())
        if len(dict_keys) > 1 and len(final_nums[dict_keys[0]]) != len(final_nums[dict_keys[-1]]):
            final_nums[dict_keys[-1]] = [final_nums[dict_keys[-1]][i] if i < len(final_nums[dict_keys[-1]]) else np.nan
                                         for i in range(0, len(final_nums[dict_keys[0]]))]
        write_df = pd.DataFrame(final_nums)
        write_df.to_excel(incoming_data['path_account_num'], header=False, index=False)
        del_pdf_files(logging, start_path)
        return {'status': 'success', 'text': f'Документы в папке {start_path.name} напечатаны'}
    except Exception as ex:  # Если ошибка
        logging.error(f"Ошибка:\n {ex}'\n'{traceback.format_exc()}")
        del_pdf_files(logging, start_path)
        return {'status': 'error', 'text': f'Ошибка при печати документов в папке {start_path.name}', 'trace': ex}


def print_docs(incoming_data: dict, current_progress, now_doc, all_doc, line_doing, line_progress, progress_value,
               event, window_check, info_value) -> dict:
    """Печать документов."""
    logging = incoming_data['logging']
    try:
        if incoming_data['package']:
            folders = [item for item in Path(incoming_data['start_path']).glob('*') if os.path.isdir(item)]
            start_folders = [Path(incoming_data['start_path'], item) for item in folders]
        else:
            start_folders = [Path(incoming_data['start_path'])]
        for start_folder in start_folders:
            logging.info(f'Бежим по папке {start_folder.name}')
            answer = folder_print(incoming_data, start_folder,
                                  line_doing, line_progress, progress_value,
                                  event, window_check, info_value)
            if answer['status'] != 'success':
                return answer
        return {'status': 'success', 'text': '', 'trace': ''}
    except BaseException as error:
        return {'status': 'error', 'text': error, 'trace': traceback.format_exc()}
