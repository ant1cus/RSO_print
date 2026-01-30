import re
import traceback

import numpy
import openpyxl
import pandas as pd
import openpyxl.styles
from natsort import natsorted
from pathlib import Path
from openpyxl.utils import get_column_letter


def create_form_27(documents: pd.DataFrame, finish_path: Path, incoming: dict):
    try:
        table_name = ['Порядковый номер', 'Дата регистрации',
                      'Номер, дата поступившего документа и гриф секретности',
                      'Откуда (от кого) поступил или кому направлен документ',
                      'Наименование или краткое содержание документа', 'Фамилия исполнителя и подразделение',
                      'экземпляров и их номера', 'листов в экземпляре', 'листов основного документа',
                      'листов приложения',
                      'Номера блока и листов черновика', 'Номера перепечатанных листов',
                      'Отметка об уничтожении черновиков', 'Кому выдан документ',
                      'Расписка в получении документа и дата', 'Номер реестра и дата',
                      'Местонахождение документа(номер дела и листа, номер акта на уничтожение и дата)',
                      'Примечание']
        table_df = pd.DataFrame({'Порядковый номер': [],
                                 'Дата регистрации': [],
                                 'Номер, дата поступившего документа и гриф секретности': [],
                                 'Откуда (от кого) поступил или кому направлен документ': [],
                                 'Наименование или краткое содержание документа': [],
                                 'Фамилия исполнителя и подразделение': [],
                                 'экземпляров и их номера': [],
                                 'листов в экземпляре': [],
                                 'листов основного документа': [],
                                 'листов приложения': [],
                                 'Номера блока и листов черновика': [],
                                 'Номера перепечатанных листов': [],
                                 'Отметка об уничтожении черновиков': [],
                                 'Кому выдан документ': [],
                                 'Расписка в получении документа и дата': [],
                                 'Номер реестра и дата': [],
                                 'Местонахождение документа(номер дела и листа, номер акта на уничтожение'
                                 ' и дата)': [],
                                 'Примечание': []})
        for document in documents.itertuples():
            index = len(table_df)
            if re.findall(r'2 экз.', document.name.lower()):
                continue
            if str(document.num_scroll) != '1':
                table_df.loc[index, 'экземпляров и их номера'] = '№' + str(document.num_scroll)
                table_df.loc[index, 'листов в экземпляре'] = document.pages
                index += 1
                table_df.loc[index] = pd.Series([numpy.nan for _ in range(0, len(table_name))], index=table_name)
                continue
            table_df.loc[index, 'Порядковый номер'] = document.footer_text
            table_df.loc[index, 'Дата регистрации'] = document.date
            table_df.loc[index, 'Номер, дата поступившего документа и гриф секретности'] = document.classified
            table_df.loc[index, 'Откуда (от кого) поступил или кому направлен документ'] = incoming['form27_firm']
            table_df.loc[index, 'Наименование или краткое содержание документа'] = document.name
            table_df.loc[index, 'Фамилия исполнителя и подразделение'] = document.executor
            if re.findall(r"заключение", document.name, re.I) and incoming['conclusion_instance']:
                table_df.loc[index, 'экземпляров и их номера'] = f"{1 + len(incoming['number_instance'])}"
            elif re.findall(r"протокол", document.name, re.I) and incoming['protocol_instance']:
                table_df.loc[index, 'экземпляров и их номера'] = f"{1 + len(incoming['number_instance'])}"
            elif re.findall(r"предписание", document.name, re.I) and incoming['prescription_instance']:
                table_df.loc[index, 'экземпляров и их номера'] = f"{1 + len(incoming['number_instance'])}"
            else:
                table_df.loc[index, 'экземпляров и их номера'] = '1'
            index += 1
            table_df.loc[index, 'экземпляров и их номера'] = '№' + str(document.num_scroll)
            table_df.loc[index, 'листов в экземпляре'] = document.pages
            index += 1
            if re.findall(r'сопроводит', document.name.lower()):
                table_df.loc[index] = pd.Series([numpy.nan for _ in range(0, len(table_name))], index=table_name)
                index += 1
                table_df.loc[index, table_name[6]] = '№2'
                table_df.loc[index, table_name[7]] = document.pages
                index += 1
            table_df.loc[index] = pd.Series([numpy.nan for _ in range(0, len(table_name))], index=table_name)
        table_df.index = pd.RangeIndex(1, 1 + len(table_df))
        table_df.to_excel(Path(finish_path, 'Форма 27.xlsx'), sheet_name='27', index=False)
        column_width = [13, 11, 10, 24, 27, 13.5, 7, 7, 7, 7, 15.1, 13, 13.1, 11.4, 16, 18.3, 23.85,
                        14.1]
        wb = openpyxl.load_workbook(Path(finish_path, 'Форма 27.xlsx'))
        ws = wb.active
        ws.insert_rows(2)
        thin = openpyxl.styles.Side(border_style="thin", color="000000")
        for el in range(1, ws.max_column + 1):
            if 6 < el < 11:
                ws.cell(2, el).value = ws.cell(1, el).value
                ws.cell(2, el).alignment = openpyxl.styles.Alignment(horizontal="center",
                                                                     vertical="center", wrap_text=True)
                ws.cell(2, el).border = openpyxl.styles.Border(top=thin, left=thin, right=thin, bottom=thin)
            else:
                if el == 11:
                    ws.merge_cells(start_row=1, end_row=1, start_column=7, end_column=10)
                    ws.cell(1, 7).value = 'Количество'
                    ws.cell(1, 7).alignment = openpyxl.styles.Alignment(horizontal="center",
                                                                        vertical="center", wrap_text=True)
                ws.merge_cells(start_row=1, end_row=2, start_column=el, end_column=el)
                ws.cell(1, el).alignment = openpyxl.styles.Alignment(horizontal="center",
                                                                     vertical="center", wrap_text=True)
        wb.save(filename=str(Path(finish_path, 'Форма 27.xlsx')))
        for el in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(el)].width = column_width[el - 1]
        flag = 0
        for row in range(3, ws.max_row + 1):
            flag += 1
            for col in range(1, ws.max_column + 1):
                if flag == 4:
                    flag = 1
                if flag == 1:
                    if col == 3 and ws.cell(row, col).value:
                        ws.cell(row, col).alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
                        ws.merge_cells(start_row=row, end_row=row + 1, start_column=col, end_column=col)
                if ws.cell(row, col).value == 0:
                    ws.cell(row, col).value = ''
                ws.cell(row, col).border = openpyxl.styles.Border(top=thin, left=thin, right=thin, bottom=thin)
        wb.save(filename=Path(finish_path, 'Форма 27.xlsx'))
        numbers = documents['footer_text'].to_numpy().tolist()
        numbers = natsorted(numbers)
        name_wb = 'Форма 27 ' + str(numbers[0]).replace('/', ',') + ' - ' + str(numbers[-1]).replace('/', ',') + '.xlsx'
        Path(finish_path, 'Форма 27.xlsx').rename(Path(finish_path, name_wb))

        return {'status': 'success', 'text': f'Форма 27 заполнена и сохранена'}
    except PermissionError as ex:
        return {'status': 'permission denied', 'text': f'Ошибка доступа при создании 27 формы', 'trace': ex}
    except BaseException as ex:
        return {'status': 'error', 'text': f'Ошибка при создании 27 формы - {ex}', 'trace': traceback.format_exc()}
